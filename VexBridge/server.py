from __future__ import annotations

import asyncio
import base64
import hmac
import inspect
import fnmatch
import io
import json
import os
import queue
import re
import shutil
import socket
import subprocess
import tempfile
import threading
import time
import urllib.request
import uuid
import zipfile
from collections import deque
from dataclasses import dataclass, field
from functools import wraps
from pathlib import Path
from typing import Any, Literal

import httpx
import psutil
from a2a.client import A2ACardResolver, ClientConfig, create_client
from a2a.helpers import new_text_message
from a2a.types import Role, SendMessageRequest
from docx import Document
from mcp.server.fastmcp import FastMCP
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.responses import JSONResponse
import uvicorn

try:
    from .brain_router import chat as brain_chat_impl, status as brain_status_impl
except ImportError:
    from brain_router import chat as brain_chat_impl, status as brain_status_impl

APP_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / "VexBridgeDC"
CONFIG_PATH = APP_DIR / "config.json"
AUDIT_PATH = APP_DIR / "audit.jsonl"
APP_DIR.mkdir(parents=True, exist_ok=True)

DEFAULT_CONFIG = {
    "allowedRoots": [str(Path.home())],
    "allowedDirectories": [str(Path.home())],
    "fileReadLineLimit": 1000,
    "fileWriteLineLimit": 30,
    "searchResultLimit": 500,
    "blockedCommands": [],
    "defaultShell": "",
    "telemetryEnabled": False,
}
if not CONFIG_PATH.exists():
    CONFIG_PATH.write_text(json.dumps(DEFAULT_CONFIG, indent=2), encoding="utf-8")

MCP_HOST = os.environ.get("VEXBRIDGE_HOST", "127.0.0.1")
MCP_PORT = int(os.environ.get("VEXBRIDGE_PORT", "8795"))
mcp = FastMCP("VexBridge", host=MCP_HOST, port=MCP_PORT, stateless_http=True, json_response=True)
RECENT_TOOL_CALLS = deque(maxlen=1000)
_NO_TRACK = {"get_recent_tool_calls", "get_usage_stats", "get_prompts", "give_feedback_to_desktop_commander"}

def tracked_tool():
    def decorate(fn):
        @wraps(fn)
        def wrapped(*args, **kwargs):
            started = time.time()
            ok = True
            result = None
            try:
                result = fn(*args, **kwargs)
                return result
            except Exception as exc:
                ok = False
                result = {"error": f"{type(exc).__name__}: {exc}"}
                raise
            finally:
                if fn.__name__ not in _NO_TRACK:
                    try:
                        arguments = dict(inspect.signature(fn).bind_partial(*args, **kwargs).arguments)
                    except Exception:
                        arguments = dict(kwargs)
                    try:
                        encoded = json.dumps(result, ensure_ascii=False, default=str)
                        output = result if len(encoded) <= 20000 else encoded[:20000] + "...[truncated]"
                    except Exception:
                        output = repr(result)[:20000]
                    RECENT_TOOL_CALLS.append({
                        "time": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(started)),
                        "tool": fn.__name__, "arguments": arguments, "output": output, "ok": ok,
                        "duration_ms": round((time.time() - started) * 1000, 2),
                    })
        return mcp.tool()(wrapped)
    return decorate

def audit(tool: str, ok: bool, detail: dict[str, Any] | None = None) -> None:
    rec = {"ts": time.time(), "tool": tool, "ok": ok, "detail": detail or {}}
    with AUDIT_PATH.open("a", encoding="utf-8") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")

def load_config() -> dict[str, Any]:
    try:
        raw = json.loads(CONFIG_PATH.read_text(encoding="utf-8-sig"))
        cfg = {**DEFAULT_CONFIG, **raw}
        if "allowedDirectories" in raw and "allowedRoots" not in raw:
            cfg["allowedRoots"] = raw["allowedDirectories"]
        elif "allowedRoots" in raw and "allowedDirectories" not in raw:
            cfg["allowedDirectories"] = raw["allowedRoots"]
        return cfg
    except Exception:
        return dict(DEFAULT_CONFIG)

def save_config(cfg: dict[str, Any]) -> None:
    if "allowedDirectories" in cfg:
        cfg["allowedRoots"] = cfg["allowedDirectories"]
    elif "allowedRoots" in cfg:
        cfg["allowedDirectories"] = cfg["allowedRoots"]
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2), encoding="utf-8")

def resolve_allowed(raw: str) -> Path:
    p = Path(os.path.expandvars(os.path.expanduser(raw))).resolve()
    cfg = load_config()
    raw_roots = cfg.get("allowedDirectories", cfg.get("allowedRoots", []))
    if raw_roots == []:
        return p
    roots = [Path(os.path.expandvars(os.path.expanduser(x))).resolve() for x in raw_roots]
    if not any(p == r or r in p.parents for r in roots):
        raise PermissionError(f"path outside allowed roots: {p}")
    return p

VEX_TOOLS_ROOT = Path(os.environ.get("VEX_TOOLS_ROOT", str(Path.home() / "Documents" / "VexNativeTools")))
ICM_BIN = Path(os.environ.get("VEX_ICM_BIN", str(VEX_TOOLS_ROOT / "ThirdParty" / "icm" / "icm.exe")))
ICM_DB = Path(os.environ.get("VEX_ICM_DB", str(Path(os.environ.get("APPDATA", str(Path.home()))) / "VexICM" / "vexnative-memory.db")))
UNLAZY_DIR = Path(os.environ.get("VEX_UNLAZY_DIR", str(VEX_TOOLS_ROOT / "ThirdParty" / "unlazy")))
A2A_BASE_URL = os.environ.get("VEX_A2A_URL", "http://127.0.0.1:8800").rstrip("/")

def _run_external(argv: list[str], timeout: int = 30, allowed_codes: set[int] | None = None) -> dict[str, Any]:
    codes = allowed_codes or {0}
    r = subprocess.run(
        argv,
        text=True,
        capture_output=True,
        timeout=timeout,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    output = (r.stdout or "") + (("\n" + r.stderr) if r.stderr else "")
    output = output.strip()
    if r.returncode not in codes:
        raise RuntimeError(f"command failed rc={r.returncode}: {output[-4000:]}")
    return {"returncode": r.returncode, "output": output}

def _icm_base() -> list[str]:
    if not ICM_BIN.exists():
        raise FileNotFoundError(f"ICM not installed: {ICM_BIN}")
    ICM_DB.parent.mkdir(parents=True, exist_ok=True)
    return [str(ICM_BIN), "--db", str(ICM_DB), "--no-embeddings"]

@tracked_tool()
def integration_status(deviceId: str | None = None) -> dict[str, Any]:
    icm_version = None
    if ICM_BIN.exists():
        try:
            icm_version = _run_external([str(ICM_BIN), "--version"], timeout=5)["output"]
        except Exception as exc:
            icm_version = f"error: {type(exc).__name__}: {exc}"
    unlazy_pin = None
    pin_path = UNLAZY_DIR / "VEX_PIN.json"
    if pin_path.exists():
        try:
            unlazy_pin = json.loads(pin_path.read_text(encoding="utf-8-sig"))
        except Exception:
            unlazy_pin = {"error": "unreadable pin file"}
    icm_http = False
    try:
        with urllib.request.urlopen("http://127.0.0.1:11435/health", timeout=2) as r:
            icm_http = int(getattr(r, "status", 0)) == 200
    except Exception:
        pass
    return {
        "icm": {
            "available": ICM_BIN.exists(),
            "version": icm_version,
            "db": str(ICM_DB),
            "httpHealthy": icm_http,
            "mode": "fts-keyword",
        },
        "unlazy": {
            "available": (UNLAZY_DIR / "scripts" / "gate-check.mjs").exists(),
            "dir": str(UNLAZY_DIR),
            "pin": unlazy_pin,
        },
        "brain": brain_status_impl(),
        "continuityAuthority": "VexContinuityVault",
    }

@tracked_tool()
def brain_status(deviceId: str | None = None) -> dict[str, Any]:
    return brain_status_impl()

@tracked_tool()
def brain_chat(
    prompt: str,
    mode: Literal["auto", "fast", "deep"] = "auto",
    system: str | None = None,
    temperature: float = 0.7,
    max_tokens: int = 4096,
    deviceId: str | None = None,
) -> dict[str, Any]:
    return brain_chat_impl(
        prompt=prompt,
        mode=mode,
        system=system,
        temperature=float(temperature),
        max_tokens=int(max_tokens),
    )

@tracked_tool()
def icm_store(
    topic: str,
    content: str,
    importance: Literal["critical", "high", "medium", "low"] = "medium",
    keywords: str | None = None,
    raw: str | None = None,
    deviceId: str | None = None,
) -> dict[str, Any]:
    argv = _icm_base() + ["store", "-t", topic, "-c", content, "-i", importance]
    if keywords:
        argv += ["-k", keywords]
    if raw:
        argv += ["-r", raw]
    result = _run_external(argv, timeout=30)
    match = re.search(r"Stored:\s*(\S+)", result["output"])
    audit("icm_store", True, {"topic": topic, "importance": importance})
    return {"ok": True, "id": match.group(1) if match else None, "output": result["output"]}

@tracked_tool()
def icm_recall(
    query: str,
    topic: str | None = None,
    limit: int = 5,
    keyword: str | None = None,
    project: str | None = None,
    deviceId: str | None = None,
) -> list[dict[str, Any]]:
    argv = _icm_base() + [
        "recall",
        query,
        "--limit",
        str(max(1, min(int(limit), 50))),
        "--format",
        "json",
    ]
    if topic:
        argv += ["--topic", topic]
    if keyword:
        argv += ["--keyword", keyword]
    if project is not None:
        argv += ["--project", project]
    result = _run_external(argv, timeout=30)
    if not result["output"]:
        return []
    data = json.loads(result["output"])
    return data if isinstance(data, list) else []

@tracked_tool()
def icm_stats(deviceId: str | None = None) -> dict[str, Any]:
    result = _run_external(_icm_base() + ["stats"], timeout=30)
    return {"db": str(ICM_DB), "mode": "fts-keyword", "output": result["output"]}

@tracked_tool()
def unlazy_status(
    gate_file: str,
    root: str | None = None,
    scope: str | None = None,
    deviceId: str | None = None,
) -> dict[str, Any]:
    gate = resolve_allowed(gate_file)
    checker = UNLAZY_DIR / "scripts" / "gate-check.mjs"
    node = shutil.which("node")
    if not checker.exists():
        raise FileNotFoundError(f"Unlazy not installed: {checker}")
    if not node:
        raise FileNotFoundError("node executable not found")
    argv = [node, str(checker), "--status"]
    if root:
        argv += ["--root", str(resolve_allowed(root))]
    if scope:
        argv += ["--scope", scope]
    argv.append(str(gate))
    result = _run_external(argv, timeout=30, allowed_codes={0, 1})
    return {"ok": result["returncode"] == 0, **result}

@tracked_tool()
def unlazy_lint(
    gate_file: str,
    strict: bool = False,
    deviceId: str | None = None,
) -> dict[str, Any]:
    gate = resolve_allowed(gate_file)
    lint = UNLAZY_DIR / "scripts" / "gate-lint.mjs"
    node = shutil.which("node")
    if not lint.exists():
        raise FileNotFoundError(f"Unlazy not installed: {lint}")
    if not node:
        raise FileNotFoundError("node executable not found")
    argv = [node, str(lint)]
    if strict:
        argv.append("--strict")
    argv.append(str(gate))
    result = _run_external(argv, timeout=30, allowed_codes={0, 1})
    return {"ok": result["returncode"] == 0, **result}


def _a2a_artifact_text(value: Any) -> list[str]:
    out: list[str] = []
    if isinstance(value, dict):
        for artifact in value.get("artifacts") or []:
            if not isinstance(artifact, dict):
                continue
            for part in artifact.get("parts") or []:
                if isinstance(part, dict) and isinstance(part.get("text"), str):
                    out.append(part["text"])
        if not out:
            for item in value.values():
                out.extend(_a2a_artifact_text(item))
    elif isinstance(value, list):
        for item in value:
            out.extend(_a2a_artifact_text(item))
    return out


def _a2a_stream_text(chunk: Any) -> str:
    found: list[str] = []
    try:
        if chunk.HasField("message"):
            found.extend(
                part.text for part in chunk.message.parts
                if getattr(part, "text", "")
            )
        if chunk.HasField("artifact_update"):
            artifact = chunk.artifact_update.artifact
            found.extend(
                part.text for part in artifact.parts
                if getattr(part, "text", "")
            )
        if chunk.HasField("task"):
            for artifact in chunk.task.artifacts:
                found.extend(
                    part.text for part in artifact.parts
                    if getattr(part, "text", "")
                )
        if chunk.HasField("status_update") and chunk.status_update.status.HasField("message"):
            found.extend(
                part.text for part in chunk.status_update.status.message.parts
                if getattr(part, "text", "")
            )
    except Exception:
        pass
    return "\n".join(found)

async def _a2a_send_async(base_url: str, message: str) -> dict[str, Any]:
    timeout = httpx.Timeout(240.0, connect=15.0)
    async with httpx.AsyncClient(timeout=timeout) as http:
        card = await A2ACardResolver(httpx_client=http, base_url=base_url).get_agent_card()
        client = await create_client(
            agent=card,
            client_config=ClientConfig(streaming=False, httpx_client=http),
        )
        try:
            request = SendMessageRequest(message=new_text_message(message, role=Role.ROLE_USER))
            events: list[dict[str, Any]] = []
            texts: list[str] = []
            async for chunk in client.send_message(request):
                extracted = _a2a_stream_text(chunk)
                if extracted:
                    texts.append(extracted)
                events.append({"value": str(chunk)})
            return {
                "agent": card.name,
                "url": base_url,
                "text": "\n".join(texts),
                "events": events,
            }
        finally:
            await client.close()

A2A_JOBS: dict[str, dict[str, Any]] = {}
A2A_JOBS_LOCK = threading.Lock()

def _a2a_job_runner(job_id: str, base_url: str, message: str) -> None:
    time.sleep(0.25)
    with A2A_JOBS_LOCK:
        A2A_JOBS[job_id].update({"status": "running", "started": time.time()})
    try:
        result = asyncio.run(_a2a_send_async(base_url, message))
        with A2A_JOBS_LOCK:
            A2A_JOBS[job_id].update({"status": "completed", "result": result, "finished": time.time()})
    except Exception as exc:
        with A2A_JOBS_LOCK:
            A2A_JOBS[job_id].update({
                "status": "failed",
                "error": f"{type(exc).__name__}: {exc}",
                "finished": time.time(),
            })

@tracked_tool()
def a2a_status(deviceId: str | None = None) -> dict[str, Any]:
    with urllib.request.urlopen(A2A_BASE_URL + "/health", timeout=5) as response:
        health = json.loads(response.read().decode("utf-8"))
    with urllib.request.urlopen(A2A_BASE_URL + "/registry", timeout=5) as response:
        registry = json.loads(response.read().decode("utf-8"))
    return {"ok": bool(health.get("ok")), "url": A2A_BASE_URL, "health": health, "registry": registry}

@tracked_tool()
def a2a_registry(deviceId: str | None = None) -> dict[str, Any]:
    with urllib.request.urlopen(A2A_BASE_URL + "/registry", timeout=5) as response:
        return json.loads(response.read().decode("utf-8"))

@tracked_tool()
def a2a_send(
    agent: Literal["coordinator", "cognition", "memory", "verification", "system", "node", "renderer", "phone", "coding"],
    message: str,
    deviceId: str | None = None,
) -> dict[str, Any]:
    suffix = "" if agent == "coordinator" else "/" + agent
    job_id = uuid.uuid4().hex
    with A2A_JOBS_LOCK:
        A2A_JOBS[job_id] = {
            "jobId": job_id,
            "agent": agent,
            "status": "submitted",
            "created": time.time(),
        }
    threading.Thread(
        target=_a2a_job_runner,
        args=(job_id, A2A_BASE_URL + suffix, message),
        daemon=True,
    ).start()
    return {"jobId": job_id, "agent": agent, "status": "submitted"}

@tracked_tool()
def a2a_result(jobId: str, deviceId: str | None = None) -> dict[str, Any]:
    with A2A_JOBS_LOCK:
        job = A2A_JOBS.get(jobId)
        if not job:
            raise KeyError(f"unknown A2A job: {jobId}")
        return dict(job)

def slice_lines(lines: list[str], offset: int = 0, length: int | None = None) -> list[str]:
    if offset < 0:
        return lines[offset:]
    lim = length if length is not None else int(load_config()["fileReadLineLimit"])
    return lines[offset:offset + lim]

def excel_values(path: Path, sheet: str | None = None, cell_range: str | None = None) -> Any:
    wb = load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[int(sheet)]] if sheet and sheet.isdigit() else wb[sheet] if sheet else wb.active
    rows = ws[cell_range] if cell_range else ws.iter_rows()
    if cell_range and not isinstance(rows, tuple):
        rows = ((rows,),)
    return [[cell.value for cell in row] for row in rows]

def docx_outline(path: Path) -> str:
    d = Document(path)
    out: list[str] = []
    for i, p in enumerate(d.paragraphs):
        if p.text:
            out.append(f"[{i}] paragraph style={p.style.name!r}: {p.text}")
    for ti, t in enumerate(d.tables):
        out.append(f"[table {ti}]")
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)

@tracked_tool()
def ping(deviceId: str | None = None) -> dict[str, Any]:
    return {"pong": True, "host": socket.gethostname(), "time": time.time()}

@tracked_tool()
def list_devices() -> list[dict[str, Any]]:
    host = socket.gethostname()
    return [{
        "id": host.lower(),
        "device_name": host,
        "status": "online",
        "auth_token": "local",
        "capabilities": {
            "app_version": "VexBridge-0.1",
            "transport": "streamable-http",
        },
    }]

@tracked_tool()
def who_am_i() -> dict[str, Any]:
    return {
        "user": os.environ.get("USERNAME") or os.environ.get("USER") or "unknown",
        "role": "local_authorized",
        "device_count": 1,
        "remote_tool_usage_percent": 0,
    }

@tracked_tool()
def get_prompts(action: Literal["get_prompt"], promptId: str, deviceId: str | None = None) -> dict[str, Any]:
    prompts = {
        "onb2_01": "Organize the Downloads folder into sensible categories and report what changed.",
        "onb2_02": "Inspect a codebase or repository and explain its structure, entry points, and important workflows.",
        "onb2_03": "Create an organized knowledge base from the selected files and folders.",
        "onb2_04": "Analyze a local data file and summarize important patterns and findings.",
        "onb2_05": "Check system health, running processes, storage, and resource pressure.",
    }
    if action != "get_prompt" or promptId not in prompts:
        raise ValueError("unsupported prompt request")
    return {"action": action, "promptId": promptId, "prompt": prompts[promptId]}

@tracked_tool()
def give_feedback_to_desktop_commander(deviceId: str | None = None) -> dict[str, Any]:
    return {"ok": True, "replacement": "VexBridge", "message": "No vendor feedback form is used by VexBridge."}

@tracked_tool()
def get_config(deviceId: str | None = None) -> dict[str, Any]:
    return load_config()

@tracked_tool()
def set_config_value(key: str, value: str | int | float | bool | list[str] | None, deviceId: str | None = None) -> dict[str, Any]:
    cfg = load_config()
    cfg[key] = value
    if key == "allowedDirectories": cfg["allowedRoots"] = value
    if key == "allowedRoots": cfg["allowedDirectories"] = value
    save_config(cfg)
    audit("set_config_value", True, {"key": key})
    return cfg

@tracked_tool()
def read_file(path: str, isUrl: bool = False, offset: int = 0, length: int = 1000,
              sheet: str | None = None, range: str | None = None, options: dict[str, Any] | None = None, deviceId: str | None = None) -> Any:
    if isUrl:
        with urllib.request.urlopen(path, timeout=30) as r:
            return r.read().decode("utf-8", errors="replace")
    p = resolve_allowed(path)
    ext = p.suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        return excel_values(p, sheet, range)
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        return {"mime": Image.MIME.get(Image.open(p).format, "application/octet-stream"),
                "base64": base64.b64encode(p.read_bytes()).decode()}
    if ext == ".pdf":
        pages = PdfReader(str(p)).pages
        selected = pages[offset:offset + length]
        return "\n\n".join(f"# Page {offset+i+1}\n{pg.extract_text() or ''}" for i, pg in enumerate(selected))
    if ext == ".docx" and offset == 0:
        return docx_outline(p)
    text = p.read_text(encoding="utf-8", errors="replace")
    return "\n".join(slice_lines(text.splitlines(), offset, length))

@tracked_tool()
def read_multiple_files(paths: list[str], deviceId: str | None = None) -> dict[str, Any]:
    out = {}
    for p in paths:
        try:
            out[p] = read_file(p)
        except Exception as e:
            out[p] = {"error": str(e)}
    return out

@tracked_tool()
def write_file(path: str, content: str, mode: str = "rewrite", deviceId: str | None = None) -> dict[str, Any]:
    p = resolve_allowed(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    ext = p.suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        data = json.loads(content)
        wb = Workbook()
        wb.remove(wb.active)
        mapping = data if isinstance(data, dict) else {"Sheet1": data}
        for name, rows in mapping.items():
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(p)
    elif ext == ".docx" and mode == "rewrite":
        d = Document()
        for line in content.splitlines():
            if line.startswith("### "): d.add_heading(line[4:], 3)
            elif line.startswith("## "): d.add_heading(line[3:], 2)
            elif line.startswith("# "): d.add_heading(line[2:], 1)
            else: d.add_paragraph(line)
        d.save(p)
    else:
        with p.open("a" if mode == "append" else "w", encoding="utf-8") as f:
            f.write(content)
    audit("write_file", True, {"path": str(p), "mode": mode})
    return {"ok": True, "path": str(p), "bytes": p.stat().st_size}

@tracked_tool()
def create_directory(path: str, deviceId: str | None = None) -> dict[str, Any]:
    p = resolve_allowed(path)
    p.mkdir(parents=True, exist_ok=True)
    return {"ok": True, "path": str(p)}

@tracked_tool()
def move_file(source: str, destination: str, deviceId: str | None = None) -> dict[str, Any]:
    s, d = resolve_allowed(source), resolve_allowed(destination)
    d.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(s), str(d))
    audit("move_file", True, {"source": str(s), "destination": str(d)})
    return {"ok": True, "destination": str(d)}

@tracked_tool()
def get_file_info(path: str, deviceId: str | None = None) -> dict[str, Any]:
    p = resolve_allowed(path)
    st = p.stat()
    info = {
        "path": str(p), "type": "directory" if p.is_dir() else "file",
        "size": st.st_size, "created": st.st_ctime, "modified": st.st_mtime,
        "mode": oct(st.st_mode)
    }
    if p.is_file() and p.suffix.lower() not in {".pdf",".docx",".xlsx",".xlsm",".png",".jpg",".jpeg",".gif",".webp"}:
        lines = p.read_text(encoding="utf-8", errors="replace").splitlines()
        info.update({"lineCount": len(lines), "lastLine": max(-1, len(lines)-1), "appendPosition": len(lines)})
    if p.suffix.lower() in {".xlsx",".xlsm"}:
        wb = load_workbook(p, read_only=True)
        info["sheets"] = [{"name": ws.title, "rowCount": ws.max_row, "colCount": ws.max_column} for ws in wb.worksheets]
    return info

@tracked_tool()
def list_directory(path: str, depth: int = 2, deviceId: str | None = None) -> list[str]:
    root = resolve_allowed(path)
    out: list[str] = []
    def walk(cur: Path, level: int) -> None:
        if level > depth: return
        try:
            items = sorted(cur.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
        except PermissionError:
            out.append(f"[DENIED] {cur.relative_to(root) if cur != root else '.'}")
            return
        for item in items[:100 if cur != root else None]:
            rel = item.relative_to(root)
            out.append(f"[DIR] {rel}" if item.is_dir() else f"[FILE] {rel}")
            if item.is_dir(): walk(item, level + 1)
        if cur != root and len(items) > 100:
            out.append(f"[WARNING] {cur.name}: {len(items)-100} items hidden")
    walk(root, 1)
    return out

@tracked_tool()
def edit_block(file_path: str, old_string: str | None = None, new_string: str | None = None,
               expected_replacements: int = 1, range: str | None = None, content: list[list[Any]] | None = None, deviceId: str | None = None) -> dict[str, Any]:
    p = resolve_allowed(file_path)
    if p.suffix.lower() in {".xlsx",".xlsm"}:
        if not range or content is None: raise ValueError("range and content required")
        wb = load_workbook(p)
        if "!" in range:
            sheet_name, rg = range.split("!", 1)
            ws = wb[sheet_name]
        else:
            ws, rg = wb.active, range
        rows = ws[rg]
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                if r_i < len(content) and c_i < len(content[r_i]):
                    cell.value = content[r_i][c_i]
        wb.save(p)
        return {"ok": True, "updated": range}
    if old_string is None or new_string is None: raise ValueError("old_string/new_string required")
    if p.suffix.lower() == ".docx":
        with zipfile.ZipFile(p, "r") as z:
            files = {n: z.read(n) for n in z.namelist()}
        names = [n for n in files if n == "word/document.xml" or n.startswith("word/header") or n.startswith("word/footer")]
        total = sum(files[n].decode("utf-8", errors="replace").count(old_string) for n in names)
        if total != expected_replacements: raise ValueError(f"expected {expected_replacements} replacements, found {total}")
        with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
            for n, data in files.items():
                if n in names:
                    txt = data.decode("utf-8", errors="replace").replace(old_string, new_string)
                    data = txt.encode("utf-8")
                z.writestr(n, data)
    else:
        txt = p.read_text(encoding="utf-8", errors="replace")
        found = txt.count(old_string)
        if found != expected_replacements: raise ValueError(f"expected {expected_replacements} replacements, found {found}")
        p.write_text(txt.replace(old_string, new_string), encoding="utf-8")
    audit("edit_block", True, {"path": str(p)})
    return {"ok": True, "replacements": expected_replacements}

def markdown_pdf(markdown: str, output: Path) -> None:
    story = []
    for block in markdown.split("\n\n"):
        if "page-break-before: always" in block:
            from reportlab.platypus import PageBreak
            story.append(PageBreak()); continue
        txt = re.sub(r"^#{1,6}\s*", "", block.strip())
        if txt:
            story.extend([Paragraph(txt.replace("\n","<br/>")), Spacer(1, 8)])
    SimpleDocTemplate(str(output), pagesize=letter).build(story)

@tracked_tool()
def write_pdf(path: str, content: Any, outputPath: str | None = None, options: dict[str, Any] | None = None, deviceId: str | None = None) -> dict[str, Any]:
    src = resolve_allowed(path)
    if isinstance(content, str):
        target = resolve_allowed(outputPath or path)
        target.parent.mkdir(parents=True, exist_ok=True)
        markdown_pdf(content, target)
        return {"ok": True, "path": str(target)}
    target = resolve_allowed(outputPath) if outputPath else None
    if target is None: raise ValueError("outputPath required for PDF modification")
    reader = PdfReader(str(src))
    pages = list(reader.pages)
    for op in content:
        if op["type"] == "delete":
            for idx in sorted(op["pageIndexes"], reverse=True):
                del pages[idx]
        elif op["type"] == "insert":
            idx = op["pageIndex"]
            if op.get("sourcePdfPath"):
                insert_pages = list(PdfReader(str(resolve_allowed(op["sourcePdfPath"]))).pages)
            else:
                with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
                    tmp = Path(tf.name)
                markdown_pdf(op.get("markdown",""), tmp)
                insert_pages = list(PdfReader(str(tmp)).pages)
                tmp.unlink(missing_ok=True)
            pages[idx:idx] = insert_pages
    writer = PdfWriter()
    for pg in pages: writer.add_page(pg)
    with target.open("wb") as f: writer.write(f)
    return {"ok": True, "path": str(target)}

@dataclass
class SearchJob:
    id: str
    results: list[dict[str, Any]] = field(default_factory=list)
    done: bool = False
    stop: bool = False
    started: float = field(default_factory=time.time)

SEARCHES: dict[str, SearchJob] = {}

def _hidden(path: Path) -> bool:
    if path.name.startswith("."): return True
    try: return bool(getattr(path.stat(), "st_file_attributes", 0) & 2)
    except Exception: return False

def search_worker(job: SearchJob, root: Path, pattern: str, search_type: str, file_pattern: str | None,
                  ignore_case: bool, literal: bool, include_hidden: bool, max_results: int,
                  timeout_ms: int | None, context_lines: int, early_termination: bool | None) -> None:
    flags = re.I if ignore_case else 0
    rx = None if literal else re.compile(pattern, flags)
    needle = pattern.lower() if ignore_case else pattern
    deadline = time.time() + timeout_ms / 1000 if timeout_ms else None
    early = (search_type == "files") if early_termination is None else early_termination
    try:
        for base, dirs, files in os.walk(root):
            if job.stop or (deadline and time.time() >= deadline): break
            if not include_hidden: dirs[:] = [d for d in dirs if not _hidden(Path(base) / d)]
            for name in files:
                if job.stop or (deadline and time.time() >= deadline): break
                p = Path(base) / name
                if not include_hidden and _hidden(p): continue
                if file_pattern and not any(fnmatch.fnmatch(name, x) for x in file_pattern.split("|")): continue
                try:
                    if search_type == "files":
                        hay = name.lower() if ignore_case else name
                        matched = needle in hay if literal else bool(rx.search(name))
                        if matched:
                            job.results.append({"path": str(p)})
                            if early and hay == needle: job.stop = True
                    else:
                        lines = p.read_text(encoding="utf-8", errors="ignore").splitlines()
                        for idx, line in enumerate(lines):
                            hay = line.lower() if ignore_case else line
                            matched = needle in hay if literal else bool(rx.search(line))
                            if matched:
                                lo=max(0,idx-context_lines); hi=min(len(lines),idx+context_lines+1)
                                job.results.append({"path":str(p),"lineNumber":idx+1,"line":line,"context":lines[lo:hi]})
                                if len(job.results) >= max_results: job.stop = True; break
                    if len(job.results) >= max_results: job.stop = True
                except Exception: pass
                if job.stop: break
            if job.stop: break
    finally:
        job.done = True

@tracked_tool()
def start_search(pattern: str, path: str, maxResults: int | None = None, includeHidden: bool = False,
                 timeout_ms: int | None = None, contextLines: int = 5, filePattern: str | None = None,
                 ignoreCase: bool = True, searchType: Literal["files", "content"] = "files",
                 earlyTermination: bool | None = None, deviceId: str | None = None,
                 literalSearch: bool = False) -> dict[str, Any]:
    root = resolve_allowed(path)
    limit = int(maxResults or load_config().get("searchResultLimit", 500))
    sid = uuid.uuid4().hex
    job = SearchJob(sid); SEARCHES[sid] = job
    threading.Thread(target=search_worker, args=(job,root,pattern,searchType,filePattern,ignoreCase,literalSearch,includeHidden,limit,timeout_ms,contextLines,earlyTermination), daemon=True).start()
    return {"sessionId": sid}

@tracked_tool()
def get_more_search_results(sessionId: str, offset: int = 0, length: int = 100, deviceId: str | None = None) -> dict[str, Any]:
    job = SEARCHES[sessionId]
    return {"results": job.results[offset:offset+length], "done": job.done, "total": len(job.results)}

@tracked_tool()
def stop_search(sessionId: str, deviceId: str | None = None) -> dict[str, Any]:
    SEARCHES[sessionId].stop = True
    return {"ok": True}

@tracked_tool()
def list_searches(deviceId: str | None = None) -> list[dict[str, Any]]:
    return [{"sessionId":k,"done":v.done,"results":len(v.results),"runtime":time.time()-v.started} for k,v in SEARCHES.items()]

@dataclass
class ProcSession:
    popen: subprocess.Popen
    started: float
    output: list[str] = field(default_factory=list)
    cursor: int = 0

PROCS: dict[int, ProcSession] = {}

def pump_output(sess: ProcSession) -> None:
    assert sess.popen.stdout is not None
    for line in iter(sess.popen.stdout.readline, ""):
        sess.output.append(line.rstrip("\r\n"))

@tracked_tool()
def start_process(timeout_ms: int, command: str, verbose_timing: bool = False, shell: str | None = None, deviceId: str | None = None) -> dict[str, Any]:
    exe = shell or None
    p = subprocess.Popen(command, shell=True if exe is None else False, executable=exe,
                         stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                         text=True, bufsize=1)
    sess = ProcSession(p, time.time()); PROCS[p.pid] = sess
    threading.Thread(target=pump_output, args=(sess,), daemon=True).start()
    deadline = time.time() + max(0, timeout_ms)/1000
    while time.time() < deadline and p.poll() is None and not sess.output:
        time.sleep(0.05)
    return {"pid": p.pid, "running": p.poll() is None, "output": sess.output[-100:]}

@tracked_tool()
def read_process_output(pid: int, timeout_ms: int = 1000, offset: int = 0, length: int = 1000,
                        verbose_timing: bool = False, deviceId: str | None = None) -> dict[str, Any]:
    s = PROCS[pid]
    if offset == 0:
        deadline = time.time() + min(timeout_ms,10000)/1000
        while time.time() < deadline and len(s.output) <= s.cursor and s.popen.poll() is None: time.sleep(0.05)
        data = s.output[s.cursor:s.cursor+length]; s.cursor += len(data)
    elif offset < 0:
        data = s.output[offset:][:length]
    else:
        data = s.output[offset:offset+length]
    return {"pid": pid, "output": data, "running": s.popen.poll() is None, "returncode": s.popen.poll()}

@tracked_tool()
def interact_with_process(pid: int, input: str, timeout_ms: int = 8000, wait_for_prompt: bool = True,
                          verbose_timing: bool = False, deviceId: str | None = None) -> dict[str, Any]:
    s = PROCS[pid]
    if s.popen.stdin is None: raise RuntimeError("stdin unavailable")
    before = len(s.output)
    s.popen.stdin.write(input + "\n"); s.popen.stdin.flush()
    deadline = time.time() + min(timeout_ms,10000)/1000
    while wait_for_prompt and time.time() < deadline and len(s.output) == before and s.popen.poll() is None: time.sleep(0.05)
    return {"pid": pid, "output": s.output[before:], "running": s.popen.poll() is None}

@tracked_tool()
def force_terminate(pid: int, deviceId: str | None = None) -> dict[str, Any]:
    s = PROCS[pid]
    s.popen.kill()
    return {"ok": True, "pid": pid}

@tracked_tool()
def list_sessions(deviceId: str | None = None) -> list[dict[str, Any]]:
    return [{"pid": pid, "running": s.popen.poll() is None, "runtime": time.time()-s.started, "lines": len(s.output)} for pid,s in PROCS.items()]

@tracked_tool()
def list_processes(deviceId: str | None = None) -> list[dict[str, Any]]:
    out = []
    for p in psutil.process_iter(["pid","name","cpu_percent","memory_info","cmdline"]):
        try:
            i=p.info; out.append({"pid":i["pid"],"name":i["name"],"cpu":i["cpu_percent"],
                "memory": getattr(i["memory_info"],"rss",None),"cmdline":i["cmdline"]})
        except (psutil.NoSuchProcess, psutil.AccessDenied):
            pass
    return out

@tracked_tool()
def kill_process(pid: int, deviceId: str | None = None) -> dict[str, Any]:
    psutil.Process(pid).kill()
    audit("kill_process", True, {"pid": pid})
    return {"ok": True, "pid": pid}

@tracked_tool()
def get_usage_stats(deviceId: str | None = None) -> dict[str, Any]:
    count = 0
    if AUDIT_PATH.exists():
        with AUDIT_PATH.open("r", encoding="utf-8") as f:
            count = sum(1 for _ in f)
    return {"localToolCallsLogged": count, "quotaPercent": 0}

@tracked_tool()
def get_recent_tool_calls(maxResults: int = 50, toolName: str | None = None, since: str | None = None, deviceId: str | None = None) -> list[dict[str, Any]]:
    rows = list(RECENT_TOOL_CALLS)
    if toolName: rows = [r for r in rows if r.get("tool") == toolName]
    if since: rows = [r for r in rows if str(r.get("time", "")) >= since]
    return rows[-max(1, min(int(maxResults), 1000)): ]

@tracked_tool()
def shutdown(deviceId: str | None = None) -> dict[str, Any]:
    def later():
        time.sleep(0.25)
        os._exit(0)
    threading.Thread(target=later, daemon=True).start()
    return {"ok": True, "message": "VexBridge shutting down"}

class BearerAuthMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        token = os.environ.get("VEXBRIDGE_TOKEN", "")
        if token:
            supplied = request.headers.get("authorization", "")
            expected = "Bearer " + token
            if not hmac.compare_digest(supplied, expected):
                return JSONResponse({"error": "unauthorized"}, status_code=401)
        return await call_next(request)

if __name__ == "__main__":
    app = mcp.streamable_http_app()
    app.add_middleware(BearerAuthMiddleware)
    uvicorn.run(app, host=MCP_HOST, port=MCP_PORT, log_level="info")
